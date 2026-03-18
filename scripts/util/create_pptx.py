"""Generate SmartGridComm benchmark presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color scheme
DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT_BLUE = RGBColor(0x00, 0x7A, 0xCC)
ACCENT_GREEN = RGBColor(0x2E, 0xA0, 0x43)
ACCENT_ORANGE = RGBColor(0xE8, 0x8D, 0x2A)
ACCENT_RED = RGBColor(0xCC, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)
VERY_LIGHT = RGBColor(0xF5, 0xF5, 0xF5)
LIGHT_BLUE_BG = RGBColor(0xE8, 0xF4, 0xFD)
TABLE_HEADER = RGBColor(0x00, 0x56, 0x8A)


def add_bg(slide, color=VERY_LIGHT):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, text, y=Inches(0.3)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), y, prs.slide_width, Inches(0.9)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT_BLUE
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_text_box(slide, text, left, top, width, height, font_size=18, bold=False, color=DARK_TEXT, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = alignment
    return tf


def add_bullet_list(slide, items, left, top, width, height, font_size=16, color=DARK_TEXT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
    return tf


def add_card(slide, title, body, left, top, width, height, accent_color=ACCENT_BLUE):
    # Card background
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    card.line.width = Pt(1)
    # Accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()
    # Title
    add_text_box(slide, title, left + Inches(0.25), top + Inches(0.1), width - Inches(0.4), Inches(0.4), font_size=18, bold=True, color=accent_color)
    # Body
    add_text_box(slide, body, left + Inches(0.25), top + Inches(0.5), width - Inches(0.4), height - Inches(0.6), font_size=14, color=MEDIUM_GRAY)


# ========== SLIDE 1: TITLE ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide, DARK_BG)

# Title
add_text_box(slide, "Benchmarking Reinforcement Learning for\nSustainable Energy Systems", Inches(0.8), Inches(1.5), Inches(11.5), Inches(2.0), font_size=40, bold=True, color=WHITE, alignment=PP_ALIGN.LEFT)
# Subtitle
add_text_box(slide, "Standard RL  |  Safe RL  |  Multi-Agent RL  |  Noise Robustness", Inches(0.8), Inches(3.5), Inches(11), Inches(0.6), font_size=22, color=ACCENT_BLUE)
# Separator line
sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.3), Inches(3), Inches(0.04))
sep.fill.solid()
sep.fill.fore_color.rgb = ACCENT_BLUE
sep.line.fill.background()
# Author info
add_text_box(slide, "Mehmet Koruturk\nVirginia Tech, RoleLab", Inches(0.8), Inches(4.6), Inches(5), Inches(0.8), font_size=20, color=LIGHT_GRAY)
# Venue
add_text_box(slide, "Target: IEEE SmartGridComm 2026  |  Deadline: June 19, 2026", Inches(0.8), Inches(5.6), Inches(8), Inches(0.5), font_size=16, color=MEDIUM_GRAY)
# Three env icons (text-based)
for i, (env, col) in enumerate([("EV Charging", ACCENT_GREEN), ("Building HVAC", ACCENT_BLUE), ("Cogeneration", ACCENT_ORANGE)]):
    x = Inches(8.5) + Inches(i * 1.6)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(5.0), Inches(1.4), Inches(1.0))
    box.fill.solid()
    box.fill.fore_color.rgb = col
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = env
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER


# ========== SLIDE 2: OVERVIEW / SCOPE ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Project Overview & Scope")

# Left: Scope description
add_text_box(slide, "Research Question", Inches(0.5), Inches(1.5), Inches(5.5), Inches(0.4), font_size=22, bold=True, color=ACCENT_BLUE)
add_text_box(slide, "How do different RL paradigms perform across diverse\nsustainable energy domains under realistic conditions?", Inches(0.5), Inches(2.0), Inches(5.5), Inches(1.0), font_size=18, color=DARK_TEXT)

# Experimental matrix
add_text_box(slide, "Experimental Matrix: 3 Environments x 3 Paradigms x 3 Noise Types", Inches(0.5), Inches(3.2), Inches(12), Inches(0.4), font_size=18, bold=True, color=DARK_TEXT)

# 3x3 grid
envs = ["EV Charging\n(54 stations, ACN-Data)", "Building HVAC\n(RC thermal model)", "Cogeneration\n(ONNX surrogate)"]
paradigms = ["Standard RL\n(PPO, SAC, TD3 via SB3)", "Safe RL\n(CMDP via OmniSafe)", "Multi-Agent RL\n(RLlib + PettingZoo)"]
colors = [ACCENT_GREEN, ACCENT_BLUE, ACCENT_ORANGE]

for j, (par, col) in enumerate(zip(paradigms, colors)):
    y = Inches(4.0) + Inches(j * 1.1)
    # Label
    add_text_box(slide, par, Inches(0.5), y, Inches(3.0), Inches(0.9), font_size=13, bold=True, color=col)
    for i, env in enumerate(envs):
        x = Inches(3.8) + Inches(i * 3.2)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.0), Inches(0.9))
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = col
        box.line.width = Pt(2)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = env
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_TEXT
        p.alignment = PP_ALIGN.CENTER


# ========== SLIDE 3: ENVIRONMENTS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Three SustainGym Environments")

env_data = [
    ("EV Charging", ACCENT_GREEN,
     "Source: ACN-Data (Caltech)\nEpisode: 288 steps (24h, 5-min)\nStations: 54\nAction: Continuous [0,32] Amps\nReward: profit - carbon - excess",
     "Obs: Dict{timestep, departures,\ndemands, MOER, forecasted_MOER}\nSafe RL cost: excess_charge * 100\nMARL: 54 agents (1 per station)"),
    ("Building HVAC", ACCENT_BLUE,
     "Source: EnergyPlus RC model\nEpisode: 288 steps (24h, 5-min)\nZones: Configurable (OfficeSmall)\nAction: Continuous [-1,1] per zone\nReward: -(energy + temp_error)",
     "Obs: Box(n+4) zone temps +\nweather + occupancy\nSafe RL cost: temp violations\nMARL: 1 agent per AC zone"),
    ("Cogeneration", ACCENT_ORANGE,
     "Source: ONNX surrogate model\nEpisode: 96 steps (1 day)\nComponents: 3 GTs + ST + aux\nAction: Mixed cont/discrete (15 keys)\nReward: -(fuel + ramp + penalties)/1e7",
     "Obs: Dict{Time, Prev_Action,\nambient, targets, prices}\nSafe RL cost: dyn_cv / 1e7\nMARL: 4 agents (GT1,GT2,GT3,ST)"),
]

for i, (name, color, left_text, right_text) in enumerate(env_data):
    y = Inches(1.5) + Inches(i * 2.0)
    # Name banner
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), y, Inches(2.2), Inches(1.8))
    banner.fill.solid()
    banner.fill.fore_color.rgb = color
    banner.line.fill.background()
    tf = banner.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = name
    p.font.size = Pt(20)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    # Left details
    add_text_box(slide, left_text, Inches(2.7), y + Inches(0.05), Inches(4.8), Inches(1.7), font_size=13, color=DARK_TEXT)
    # Right details
    add_text_box(slide, right_text, Inches(7.7), y + Inches(0.05), Inches(5.3), Inches(1.7), font_size=13, color=MEDIUM_GRAY)


# ========== SLIDE 4: NOISE ROBUSTNESS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Noise Robustness Framework")

add_text_box(slide, "Three Independent Noise Channels", Inches(0.5), Inches(1.5), Inches(12), Inches(0.4), font_size=22, bold=True, color=DARK_TEXT)

noise_data = [
    ("Observation Noise", "Applied after env step, before agent observes\nSimulates sensor uncertainty, forecast errors\nScales: 0.0, 0.05, 0.10, 0.20, 0.40", ACCENT_BLUE, "agent sees noisy state"),
    ("Action Noise", "Applied before env processes action\nSimulates actuator imprecision, communication errors\nScales: 0.0, 0.05, 0.10, 0.20, 0.40", ACCENT_GREEN, "env receives noisy action"),
    ("Environment Noise", "Applied in dynamics/reward computation\nSimulates real-world stochasticity, model mismatch\nScales: 0.0, 0.05, 0.10, 0.20, 0.40", ACCENT_ORANGE, "dynamics are stochastic"),
]

for i, (title, desc, color, tag) in enumerate(noise_data):
    y = Inches(2.2) + Inches(i * 1.6)
    add_card(slide, title, desc, Inches(0.5), y, Inches(7.5), Inches(1.4), color)
    # Tag
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), y + Inches(0.4), Inches(4.0), Inches(0.6))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = color
    tag_box.line.fill.background()
    tf = tag_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = tag
    p.font.size = Pt(16)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

# Key insight box
add_text_box(slide, "Key Question: Which noise type degrades performance most? (Hypothesis: env noise > obs noise > action noise)", Inches(0.5), Inches(6.5), Inches(12), Inches(0.5), font_size=16, bold=True, color=ACCENT_RED)


# ========== SLIDE 5: FEASIBILITY ASSESSMENT ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Publication Feasibility Assessment")

# Strengths
add_text_box(slide, "STRENGTHS", Inches(0.5), Inches(1.5), Inches(5.8), Inches(0.4), font_size=20, bold=True, color=ACCENT_GREEN)
strengths = [
    "3 diverse environments spanning EV, building, cogen domains",
    "3 evaluation axes (standard RL + safe RL + MARL) — rare in literature",
    "Structured noise decomposition (obs / action / env) — novel",
    "Production-quality codebase (SB3 + OmniSafe + RLlib)",
    "Realistic env models (ACN-Data, EnergyPlus RC, ONNX surrogate)",
    "Extensive EVCharging results (15+ safe RL algorithms tested)",
]
add_bullet_list(slide, [f"+ {s}" for s in strengths], Inches(0.5), Inches(2.0), Inches(5.8), Inches(3.5), font_size=14, color=DARK_TEXT)

# Weaknesses
add_text_box(slide, "GAPS TO ADDRESS", Inches(6.8), Inches(1.5), Inches(6.0), Inches(0.4), font_size=20, bold=True, color=ACCENT_RED)
weaknesses = [
    "No non-RL baselines (rule-based, MPC) — CRITICAL",
    "Inconsistent experimental completeness across envs",
    "No documented statistical rigor (seeds, CIs)",
    "Missing domain-specific metrics beyond reward",
    "Env noise experiments incomplete",
    "No cross-axis (StdRL vs SafeRL vs MARL) comparison",
]
add_bullet_list(slide, [f"- {w}" for w in weaknesses], Inches(6.8), Inches(2.0), Inches(6.0), Inches(3.5), font_size=14, color=DARK_TEXT)

# Verdict
verdict = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(5.8), Inches(12.3), Inches(1.2))
verdict.fill.solid()
verdict.fill.fore_color.rgb = LIGHT_BLUE_BG
verdict.line.color.rgb = ACCENT_BLUE
verdict.line.width = Pt(2)
tf = verdict.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.3)
p = tf.paragraphs[0]
p.text = "VERDICT: Publishable at SmartGridComm with targeted improvements. Scope is right; execution needs completing."
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.LEFT


# ========== SLIDE 6: LITERATURE POSITIONING ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Literature Positioning & Gap")

add_text_box(slide, "Existing Benchmarks", Inches(0.5), Inches(1.5), Inches(6), Inches(0.4), font_size=20, bold=True, color=ACCENT_BLUE)

benchmarks = [
    ("SustainGym (NeurIPS'23)", "3 envs, basic RL baselines only", "We extend with safe RL, MARL, noise"),
    ("CityLearn (BuildSys'19)", "Buildings only, MARL focus", "We cover 3 domains, add safe RL"),
    ("BOPTEST (JBPS'21)", "Buildings only, includes MPC", "We add RL benchmarking protocol"),
    ("Sinergym (BuildSys'21)", "EnergyPlus RL, single domain", "We are multi-domain, multi-paradigm"),
    ("Grid2Op / L2RPN", "Transmission grid, different domain", "Complementary; different scale"),
]

for i, (name, scope, gap) in enumerate(benchmarks):
    y = Inches(2.1) + Inches(i * 0.75)
    add_text_box(slide, name, Inches(0.5), y, Inches(3.0), Inches(0.6), font_size=14, bold=True, color=DARK_TEXT)
    add_text_box(slide, scope, Inches(3.8), y, Inches(4.0), Inches(0.6), font_size=13, color=MEDIUM_GRAY)
    add_text_box(slide, gap, Inches(8.2), y, Inches(4.8), Inches(0.6), font_size=13, color=ACCENT_GREEN)

# Gap summary
add_text_box(slide, "Our Unique Position", Inches(0.5), Inches(5.8), Inches(12), Inches(0.4), font_size=20, bold=True, color=ACCENT_BLUE)
gaps = [
    "No existing benchmark covers EV + Building + Cogen under a unified protocol",
    "No systematic safe RL benchmark across multiple energy domains",
    "No structured obs/action/env noise decomposition for energy RL",
    "No paper compares standard RL, safe RL, and MARL on the same environments",
]
add_bullet_list(slide, gaps, Inches(0.5), Inches(6.3), Inches(12), Inches(1.5), font_size=15, color=DARK_TEXT)


# ========== SLIDE 7: GAP ANALYSIS - WHAT'S MISSING ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Gap Analysis: What Must Be Done")

items = [
    ("P0: Non-RL Baselines", "Rule-based + optimization baseline per env. Answers 'why RL?'", ACCENT_RED, "CRITICAL"),
    ("P0: Statistical Rigor", "5+ seeds per experiment, bootstrap 95% CIs", ACCENT_RED, "CRITICAL"),
    ("P0: Domain Metrics", "Energy cost ($), CO2 (kg), constraint violation (%),\ncomfort index, peak demand (kW)", ACCENT_RED, "CRITICAL"),
    ("P0: Complete All Envs", "Building & Cogen have gaps in noise sweeps,\nsafe RL, and MARL training", ACCENT_RED, "CRITICAL"),
    ("P1: Cross-Axis Comparison", "StdRL vs SafeRL vs MARL on each env — the\ncentral deliverable table", ACCENT_ORANGE, "HIGH"),
    ("P1: Distribution Shift Test", "Train clean, test noisy — most practical\nrobustness evaluation", ACCENT_ORANGE, "HIGH"),
    ("P2: Noise Decomposition Ablation", "Show obs vs action vs env noise have different\nimpact profiles", ACCENT_BLUE, "MEDIUM"),
]

for i, (title, desc, color, priority) in enumerate(items):
    y = Inches(1.5) + Inches(i * 0.82)
    # Priority badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), y + Inches(0.05), Inches(1.2), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = color
    badge.line.fill.background()
    tf = badge.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = priority
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    # Content
    add_text_box(slide, title, Inches(1.7), y, Inches(3.5), Inches(0.35), font_size=15, bold=True, color=DARK_TEXT)
    add_text_box(slide, desc, Inches(5.3), y, Inches(7.5), Inches(0.7), font_size=13, color=MEDIUM_GRAY)


# ========== SLIDE 8: EXPERIMENT STATUS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Current Experiment Completion Status")

# Table-like layout
headers = ["Experiment", "EVCharging", "Building", "Cogen"]
rows = [
    ["Std RL (PPO/SAC/TD3)", "Complete", "Partial", "Partial"],
    ["Obs Noise Sweep", "Complete", "Partial", "Partial"],
    ["Action Noise Sweep", "Complete", "Partial", "Missing"],
    ["Env Noise Sweep", "In Progress", "Missing", "Missing"],
    ["Safe RL (OmniSafe)", "Extensive", "Early", "Incomplete"],
    ["MARL (RLlib)", "APPO only", "SAC only", "PPO (100 iter)"],
    ["Non-RL Baselines", "Missing", "Missing", "Missing"],
    ["Multi-seed (5+)", "Unknown", "Unknown", "Unknown"],
    ["Domain Metrics", "Partial", "Missing", "Missing"],
]

status_colors = {
    "Complete": ACCENT_GREEN, "Extensive": ACCENT_GREEN,
    "Partial": ACCENT_ORANGE, "In Progress": ACCENT_ORANGE,
    "Early": ACCENT_ORANGE, "APPO only": ACCENT_ORANGE,
    "SAC only": ACCENT_ORANGE, "PPO (100 iter)": ACCENT_ORANGE,
    "Incomplete": ACCENT_ORANGE,
    "Missing": ACCENT_RED, "Unknown": MEDIUM_GRAY,
}

y_start = Inches(1.6)
col_widths = [Inches(3.0), Inches(2.8), Inches(2.8), Inches(2.8)]
col_starts = [Inches(0.8)]
for w in col_widths[:-1]:
    col_starts.append(col_starts[-1] + w + Inches(0.2))

# Headers
for j, (h, x, w) in enumerate(zip(headers, col_starts, col_widths)):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y_start, w, Inches(0.5))
    box.fill.solid()
    box.fill.fore_color.rgb = TABLE_HEADER
    box.line.fill.background()
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = h
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

# Rows
for i, row in enumerate(rows):
    y = y_start + Inches(0.55) + Inches(i * 0.55)
    bg_color = WHITE if i % 2 == 0 else VERY_LIGHT
    for j, (cell, x, w) in enumerate(zip(row, col_starts, col_widths)):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.5))
        box.fill.solid()
        box.fill.fore_color.rgb = bg_color
        box.line.color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        box.line.width = Pt(0.5)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = cell
        p.font.size = Pt(13)
        p.font.bold = (j == 0)
        if j == 0:
            p.font.color.rgb = DARK_TEXT
        else:
            p.font.color.rgb = status_colors.get(cell, DARK_TEXT)
            p.font.bold = True
        p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
        if j == 0:
            tf.margin_left = Inches(0.15)


# ========== SLIDE 9: PROPOSED CONTRIBUTIONS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Proposed Contributions")

contributions = [
    ("C1", "Unified Multi-Paradigm Evaluation Protocol",
     "First benchmark comparing standard RL, safe RL, and MARL\nacross multiple energy domains with standardized metrics", ACCENT_BLUE),
    ("C2", "Systematic Noise Robustness Analysis",
     "Decomposed obs/action/env noise evaluation revealing which\nnoise type impacts energy RL most", ACCENT_GREEN),
    ("C3", "First Safe RL Benchmark for EV Charging & Cogen",
     "CMDP formulation with OmniSafe showing constraint\nviolation vs. reward tradeoff (Pareto analysis)", ACCENT_ORANGE),
    ("C4", "Cross-Domain MARL Comparison",
     "Single-agent vs multi-agent across EV (54 agents),\nbuilding (per-zone), cogen (4 agents)", ACCENT_RED),
    ("C5", "Open-Source Benchmark Suite",
     "Reproducible training, evaluation, and reporting scripts\nwith documented protocol and code release", MEDIUM_GRAY),
]

for i, (cid, title, desc, color) in enumerate(contributions):
    y = Inches(1.5) + Inches(i * 1.15)
    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), y + Inches(0.1), Inches(0.6), Inches(0.6))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    tf = circle.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = cid
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    # Title
    add_text_box(slide, title, Inches(1.3), y, Inches(5.0), Inches(0.4), font_size=18, bold=True, color=DARK_TEXT)
    # Description
    add_text_box(slide, desc, Inches(1.3), y + Inches(0.4), Inches(11.0), Inches(0.7), font_size=14, color=MEDIUM_GRAY)


# ========== SLIDE 10: NOVELTY STRENGTHENING ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Novelty Assessment & Strengthening")

add_text_box(slide, "Current Novelty: MODERATE", Inches(0.5), Inches(1.5), Inches(6), Inches(0.4), font_size=22, bold=True, color=ACCENT_ORANGE)
add_text_box(slide, "Environments are published (SustainGym NeurIPS'23). Algorithms are standard.\nNovelty must come from the benchmark protocol and insights it generates.", Inches(0.5), Inches(2.0), Inches(12), Inches(0.7), font_size=16, color=MEDIUM_GRAY)

add_text_box(slide, "If We Add...                                                            We Can Claim...", Inches(0.5), Inches(3.0), Inches(12), Inches(0.4), font_size=18, bold=True, color=DARK_TEXT)

strengthening = [
    ("Non-RL baselines (rule-based + LP/MPC)", '"RL outperforms optimization under uncertainty"'),
    ("Distribution shift experiments", '"Policy robustness analysis under deployment conditions"'),
    ("Cross-paradigm comparison table", '"Safe RL reduces violations X% at Y% reward cost"'),
    ("Open-source benchmark release", '"Reproducible benchmark suite as community artifact"'),
    ("Pareto analysis for safe RL", '"Tunable safety-reward tradeoff characterization"'),
]

for i, (add, claim) in enumerate(strengthening):
    y = Inches(3.6) + Inches(i * 0.75)
    # "Add" box
    add_card(slide, "", add, Inches(0.5), y, Inches(5.5), Inches(0.6), ACCENT_BLUE)
    # Arrow
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.2), y + Inches(0.1), Inches(0.5), Inches(0.4))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = ACCENT_GREEN
    arrow.line.fill.background()
    # "Claim" box
    add_card(slide, "", claim, Inches(6.9), y, Inches(6.0), Inches(0.6), ACCENT_GREEN)


# ========== SLIDE 11: BASELINES NEEDED ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Required Baselines")

add_text_box(slide, 'Every benchmark must answer: "Why use RL at all?"', Inches(0.5), Inches(1.5), Inches(12), Inches(0.4), font_size=20, bold=True, color=ACCENT_RED)

baseline_data = [
    ("Do-Nothing / Fixed", ACCENT_GREEN, "Trivial",
     "EV: Charge at max rate\nBuilding: Fixed thermostat setpoint\nCogen: Nominal operating point"),
    ("Rule-Based / Heuristic", ACCENT_BLUE, "Low",
     "EV: LLLF or proportional to demand\nBuilding: Deadband controller (+/- 1C)\nCogen: Load-following proportional"),
    ("Optimization (LP/MPC)", ACCENT_ORANGE, "Medium",
     "EV: LP for min-cost scheduling\nBuilding: 1-step MPC with RC model\nCogen: LP for fuel minimization"),
]

for i, (name, color, effort, desc) in enumerate(baseline_data):
    y = Inches(2.3) + Inches(i * 1.7)
    # Header
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), y, Inches(3.5), Inches(0.6))
    header.fill.solid()
    header.fill.fore_color.rgb = color
    header.line.fill.background()
    tf = header.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = f"{name}  [Effort: {effort}]"
    p.font.size = Pt(16)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    # Description
    add_text_box(slide, desc, Inches(4.3), y, Inches(8.5), Inches(1.4), font_size=15, color=DARK_TEXT)

add_text_box(slide, "Minimum: Do-Nothing + Rule-Based for all 3 envs (feasible before deadline)", Inches(0.5), Inches(6.8), Inches(12), Inches(0.4), font_size=16, bold=True, color=ACCENT_BLUE)


# ========== SLIDE 12: EVALUATION PROTOCOL ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Standardized Evaluation Protocol")

# Seeds
add_card(slide, "Statistical Rigor", "5 seeds per experiment (42-46)\nBootstrap 95% confidence intervals\nInterquartile mean (per Agarwal et al. NeurIPS'21)", Inches(0.3), Inches(1.5), Inches(6.2), Inches(1.5), ACCENT_BLUE)

# Metrics
add_card(slide, "Domain-Specific Metrics", "Energy cost ($)  |  CO2 emissions (kg)\nConstraint violation rate (%)  |  Peak demand (kW)\nComfort / service quality index", Inches(6.8), Inches(1.5), Inches(6.2), Inches(1.5), ACCENT_GREEN)

# Protocol
add_card(slide, "Noise Robustness Protocol", "For each {env x algo x noise_type x noise_level}:\n  5 seeds -> mean +/- 95% CI\nDistribution shift: train clean, test at noise={0.05,0.10,0.20,0.40}", Inches(0.3), Inches(3.3), Inches(6.2), Inches(1.5), ACCENT_ORANGE)

# Reporting
add_card(slide, "Reporting Standards", "Training curves: shaded IQR across seeds\nBar charts with error bars for final performance\nPareto fronts for safe RL (reward vs cost)", Inches(6.8), Inches(3.3), Inches(6.2), Inches(1.5), ACCENT_RED)

# Figures/Tables
add_text_box(slide, "Required Paper Deliverables", Inches(0.5), Inches(5.2), Inches(12), Inches(0.4), font_size=18, bold=True, color=DARK_TEXT)

deliverables = [
    "Fig 1: Taxonomy diagram (3 envs x 3 paradigms x noise axes)",
    "Fig 2: Training curves per env (shaded IQR)",
    "Fig 3: Noise degradation curves (reward vs noise level by type)",
    "Fig 4: Safe RL Pareto fronts (reward vs constraint cost)",
    "Fig 5: MARL vs single-agent comparison bars",
    "Table I: Env specs  |  Table II: Main results  |  Table III: Domain metrics  |  Table IV: Robustness",
]
add_bullet_list(slide, deliverables, Inches(0.5), Inches(5.7), Inches(12), Inches(2.0), font_size=14, color=DARK_TEXT)


# ========== SLIDE 13: TIMELINE ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Timeline to Submission (June 19, 2026)")

phases = [
    ("Phase 1: Critical Fixes", "Mar 3 - Apr 15", ACCENT_RED,
     ["Implement Do-Nothing + Rule-Based baselines (all envs)",
      "Run all experiments with 5 seeds",
      "Add domain-metric extraction to testing pipeline",
      "Complete Building & Cogen std RL training"]),
    ("Phase 2: Core Experiments", "Apr 15 - May 15", ACCENT_ORANGE,
     ["Safe RL for Building & Cogen (3 algos x 5 seeds)",
      "MARL training with sufficient iterations",
      "Distribution shift experiments",
      "Cross-paradigm comparison (StdRL vs SafeRL vs MARL)"]),
    ("Phase 3: Analysis & Writing", "May 15 - Jun 12", ACCENT_BLUE,
     ["Generate all figures and tables",
      "Write 6-page IEEE format paper",
      "Internal review and revision",
      "Code cleanup for public release"]),
    ("Phase 4: Submit", "Jun 12 - Jun 19", ACCENT_GREEN,
     ["Final proofreading", "Supplementary materials", "Submit!"]),
]

for i, (title, dates, color, tasks) in enumerate(phases):
    y = Inches(1.5) + Inches(i * 1.45)
    # Phase header
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), y, Inches(3.5), Inches(0.5))
    header.fill.solid()
    header.fill.fore_color.rgb = color
    header.line.fill.background()
    tf = header.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = f"{title}  ({dates})"
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    # Tasks
    for j, task in enumerate(tasks):
        add_text_box(slide, f"  {task}", Inches(4.0) + Inches(j % 2 * 4.5), y + Inches(0.05 + (j // 2) * 0.45), Inches(4.5), Inches(0.4), font_size=13, color=DARK_TEXT)


# ========== SLIDE 14: PAPER STRUCTURE ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Recommended Paper Structure (6-Page IEEE)")

sections = [
    ("I. Introduction", "0.75 pg", "Problem statement, motivation,\n5 contribution bullet points"),
    ("II. Related Work", "0.75 pg", "RL for energy, safe RL, MARL,\nbenchmarks with gap identification"),
    ("III. Environments &\n    Problem Formulation", "1.0 pg", "Table I (env specs), noise model\nformalization, CMDP formulation"),
    ("IV. Experimental Setup", "0.75 pg", "Algorithms, hyperparameters, baselines,\nevaluation protocol, seeds"),
    ("V. Results & Analysis", "2.0 pg", "Main results (Tables II-III),\nnoise robustness, safe RL, MARL"),
    ("VI. Discussion &\n     Conclusion", "0.75 pg", "Key insights, limitations,\nfuture work, code release"),
]

for i, (section, pages, content) in enumerate(sections):
    y = Inches(1.5) + Inches(i * 0.95)
    # Section name
    add_text_box(slide, section, Inches(0.5), y, Inches(3.5), Inches(0.8), font_size=16, bold=True, color=ACCENT_BLUE)
    # Pages
    add_text_box(slide, pages, Inches(4.2), y + Inches(0.1), Inches(1.0), Inches(0.4), font_size=14, bold=True, color=ACCENT_ORANGE)
    # Content
    add_text_box(slide, content, Inches(5.3), y, Inches(7.5), Inches(0.8), font_size=14, color=DARK_TEXT)


# ========== SLIDE 15: KEY QUESTIONS ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "Key Research Questions & Expected Findings")

questions = [
    ("Q1: Which algorithm wins?", "Likely SAC for off-policy efficiency,\nPPO for stability under noise", ACCENT_BLUE),
    ("Q2: Is safe RL worth the reward cost?", "Quantify Pareto tradeoff:\nX% fewer violations at Y% reward reduction", ACCENT_GREEN),
    ("Q3: Does MARL improve over single-agent?", "Likely environment-dependent:\nEV charging may benefit, cogen may not", ACCENT_ORANGE),
    ("Q4: Which noise type matters most?", "Hypothesis: env noise (distribution shift)\ndominates obs and action noise", ACCENT_RED),
    ("Q5: Does RL beat non-RL baselines?", "RL likely beats rule-based but may\nnot beat MPC in low-noise settings", MEDIUM_GRAY),
]

for i, (q, expected, color) in enumerate(questions):
    y = Inches(1.5) + Inches(i * 1.15)
    # Question
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), y, Inches(5.5), Inches(0.9))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = q
    p.font.size = Pt(16)
    p.font.color.rgb = WHITE
    p.font.bold = True
    # Expected finding
    add_text_box(slide, expected, Inches(6.2), y, Inches(6.8), Inches(0.9), font_size=15, color=DARK_TEXT)


# ========== SLIDE 16: SUMMARY ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK_BG)

add_text_box(slide, "Summary & Next Steps", Inches(0.8), Inches(0.8), Inches(11), Inches(0.6), font_size=36, bold=True, color=WHITE)

sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(3), Inches(0.04))
sep.fill.solid()
sep.fill.fore_color.rgb = ACCENT_BLUE
sep.line.fill.background()

summary_items = [
    "The benchmark scope (3 envs x 3 paradigms) is strong and fits SmartGridComm",
    "Critical gaps: non-RL baselines, statistical rigor, domain metrics, complete experiments",
    "Novelty is moderate but can be strengthened with protocol + insights framing",
    "Timeline is tight but feasible: 3.5 months to June 19 deadline",
    "Priority: baselines > multi-seed > complete experiments > write paper",
]

for i, item in enumerate(summary_items):
    y = Inches(2.0) + Inches(i * 0.7)
    # Bullet
    bullet = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), y + Inches(0.08), Inches(0.2), Inches(0.2))
    bullet.fill.solid()
    bullet.fill.fore_color.rgb = ACCENT_BLUE
    bullet.line.fill.background()
    add_text_box(slide, item, Inches(1.2), y, Inches(11), Inches(0.5), font_size=18, color=WHITE)

add_text_box(slide, "Immediate Actions", Inches(0.8), Inches(5.5), Inches(5), Inches(0.4), font_size=22, bold=True, color=ACCENT_ORANGE)

actions = [
    "1. Implement Do-Nothing + Rule-Based baselines this week",
    "2. Set up 5-seed runs for all completed experiments on ARC",
    "3. Add domain-metric extraction to stdrl_testing.py",
    "4. Queue remaining Building & Cogen experiments on ARC",
]
add_bullet_list(slide, actions, Inches(0.8), Inches(6.0), Inches(11), Inches(1.5), font_size=16, color=LIGHT_GRAY)


# ========== SAVE ==========
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "smartgridcomm_benchmark_presentation.pptx")
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
